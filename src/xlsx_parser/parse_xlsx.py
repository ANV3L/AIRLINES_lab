import pandas as pd

from pathlib import Path
import argparse

import clean_up

from openpyxl import load_workbook

def parse_file_v2(file_path, boarding_pass_data):
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            # Читаем только первые 13 строк и 8 столбцов
            rows = list(ws.iter_rows(min_row=1, max_row=13, max_col=8, values_only=True))
            if len(rows) < 13:
                continue  # пропускаем листы, где меньше 13 строк

            def get(row_idx, col_idx):
                try:
                    return rows[row_idx][col_idx]
                except (IndexError, TypeError):
                    return None

            sequence       = get(0, 7)
            name           = get(2, 1)
            seat_class     = get(2, 7)
            flight_number  = get(4, 0)
            from_city      = get(4, 3)
            to_city        = get(4, 7)
            from_airport   = get(6, 3)
            to_airport     = get(6, 7)
            date_departure = get(8, 0)
            time_departure = get(8, 2)
            pnr_code       = get(12, 1)
            e_ticket       = get(12, 4)
            bonus_card     = get(2, 5)

            boarding_pass_data["Очередь_регистрации"].append(sequence)
            boarding_pass_data["Имя_пассажира"].append(name)
            boarding_pass_data["Класс_обслуживания"].append(seat_class)
            boarding_pass_data["Номер_рейса"].append(flight_number)
            boarding_pass_data["Город_вылета"].append(from_city)
            boarding_pass_data["Город_прилета"].append(to_city)
            boarding_pass_data["Код_аэропорта_вылета"].append(from_airport)
            boarding_pass_data["Код_аэропорта_прилета"].append(to_airport)
            boarding_pass_data["Дата_вылета"].append(date_departure)
            boarding_pass_data["Время_вылета"].append(time_departure)
            boarding_pass_data["PNR_бронь"].append(pnr_code)
            boarding_pass_data["Номер_эл_билета"].append(e_ticket)
            boarding_pass_data["Бонусная_карта"].append(bonus_card)
    finally:
        wb.close()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--folder', type=str, default='data', help='Путь к папке с файлами')

    args = parser.parse_args()

    folder = Path(args.folder)
    output_file_path = "boarding_pass_raw.csv"

    files = sorted([f for f in folder.iterdir() if f.is_file()])

    boarding_pass_data = {
        "Очередь_регистрации": [],
        "Имя_пассажира": [],
        "Класс_обслуживания": [],
        "Номер_рейса": [],
        "Город_вылета": [],
        "Город_прилета": [],
        "Код_аэропорта_вылета": [],
        "Код_аэропорта_прилета": [],
        "Дата_вылета": [],
        "Время_вылета": [],
        "PNR_бронь": [],
        "Номер_эл_билета": [],
        "Бонусная_карта": [],
    }

    print("Файлов:", len(files))
    count = 0
    for file_path in files:
        print("Обработано:", count/len(files)*100, "%")
        parse_file_v2(file_path, boarding_pass_data)
        count += 1
    print("Парсинг завершен!\nСохраняю...")

    df = pd.DataFrame(boarding_pass_data)
    df.to_csv(output_file_path, index=False, sep="|")

    print("Сохранено в", output_file_path)

    clean_up.make_clean_csv(df)
if __name__ == "__main__":
    main()